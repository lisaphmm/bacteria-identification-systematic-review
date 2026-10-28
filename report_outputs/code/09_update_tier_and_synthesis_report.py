"""Update the current report using the reviewed tier decisions and saved analyses.

Creates one immutable pre-edit backup; reruns use that backup to avoid duplicate sections.
Run 08_summarise_evidence_tiers.py first. Does not rerun the legacy report builder.
"""
import copy, csv, json, shutil, re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parents[2]
REPORT=ROOT/'protocol/Systematic Review.docx'
BACKUP=ROOT/'protocol/Systematic Review (backup before tier and synthesis 2026-10-28).docx'
if not BACKUP.exists():shutil.copy2(REPORT,BACKUP)
d=Document(BACKUP)
# Make room for the inserted table/figure in order of first appearance.
for node in d._element.iter(qn('w:t')):
    s=node.text or ''
    s=re.sub(r'\bFigure ([56])([a-c]?)\b',lambda m:'Figure '+str(int(m.group(1))+1)+m.group(2),s)
    s=re.sub(r'\bTable 3\.([567])\b',lambda m:'Table 3.'+str(int(m.group(1))+1),s)
    node.text=s
TABLES=ROOT/'report_outputs/tables'
summary=json.loads((ROOT/'results/final/tier_review/analysis_summary.json').read_text())
# Narrative values below were checked for this source-review snapshot. Fail loudly
# if decisions change, rather than silently inserting stale manuscript results.
assert summary['tier_counts']=={'II':98,'I':75,'III':33}
assert summary['confidence_counts']=={'high':161,'medium':35,'low':10}
assert summary['n_accuracy']==126
dist=pd.read_csv(TABLES/'evidence_tier_distribution.csv')
acc=pd.read_csv(TABLES/'evidence_tier_accuracy_summary.csv')

def find(prefix):
    return next(p for p in d.paragraphs if p.text.strip().startswith(prefix))
def replace(p,text):
    for child in list(p._p):
        if child.tag!=qn('w:pPr'):p._p.remove(child)
    p.add_run(text)
    return p
def insert_after(el,text,style='Normal'):
    node=OxmlElement('w:p');el.addnext(node)
    p=Paragraph(node,d._body);p.style=d.styles[style];p.add_run(text)
    return p
def replace_between(start,end,texts):
    a=find(start)._p;b=find(end)._p
    node=a.getnext()
    while node is not None and node is not b:
        nxt=node.getnext();node.getparent().remove(node);node=nxt
    for text,style in texts:a=insert_after(a,text,style)._p
def normal(text):return (text,'Normal')
def rounded2(value):return str(Decimal(str(value)).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP))
def add_table(after,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    for ci,w in enumerate(widths):t.columns[ci].width=Inches(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):cells[i].text=str(val)
    for rown,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci])
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                for r in p.runs:r.font.size=Pt(8);r.bold=(rown==0)
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
    prop=t.rows[0]._tr.get_or_add_trPr();prop.append(OxmlElement('w:tblHeader'))
    after.addnext(t._tbl)
    return t._tbl

# Replace the unperformed dual-LLM/human tiering account with what was actually done.
replace(find('Tiers were assigned from the extracted fields'),
 'On 28 October 2026, one AI reviewer (Codex) individually reviewed acquisition, dataset provenance and main evaluation passages in the local full text of each of the 206 included reports, using the extracted fields as a navigation aid. Tier decisions were made by source-based reading, not by a scripted classifier. The tier-v1 definitions in evidence_tier_instructions.txt were retained; the present full-text review supersedes its suggested extracted-field-only, dual-model workflow. Neither a second independent tiering model nor human consensus adjudication was performed in this pass. Supplementary Table S2 records the tier, confidence, supporting rationale, PDF page or section, unresolved information, reviewer identity and local source-file hash for every report. The review covered relevant passages, rather than verifying every page or supplementary file.')
replace(find('We counted studies per tier'),
 'We retained the highest tier supported by a main evaluation. Original datasets released publicly by their own creators were distinguished from secondary reuse by another group. Deliberately spiked samples and reference-strain experiments remained Tier II under decision rule 4, even when a custom imaging platform was used; a prototype label, processing speed or proposed deployment did not establish live use. Where the prototype clause and reference-strain rule overlapped, the rationale and uncertainty were recorded. Missing provenance led to the lower plausible tier with low confidence. We reported tier frequencies, Wilson 95% confidence intervals, agar-plate and model-family cross-tabulations, and the distributions of as-extracted headline accuracy. These distributions were interpreted descriptively because tasks and validation units differed.')

quadas_paragraphs=[
 'A review-specific adaptation of QUADAS-2 will assess risk of bias separately from the realism tiers [65]. The sampling unit is the patient specimen, isolate, plate, slide, image or time-lapse experiment, as appropriate to the task; crops, colonies and repeated frames nested within a source unit are not automatically independent samples. For colony identification, the target state is the organism identity at the stated taxonomic level, and correct identification means agreement with that identity. The index test comprises the complete AI/ML pipeline, including acquisition-dependent preprocessing, feature extraction, model fitting and any decision threshold.',
 'Reference standards must match the outcome. MALDI-TOF, validated biochemical identification, sequencing or a stated microbiological identification procedure can establish organism identity. Expert annotations and manual counts can establish object locations, boundaries or colony numbers, but cannot alone verify species identity. For microscopy detection, the reference may be expert-confirmed bacterial presence; clinical infection requires an appropriate clinical or microbiological reference. Counting, segmentation and antimicrobial-susceptibility studies will therefore retain task-specific reference standards and outcomes, rather than being treated as interchangeable diagnostic-identification studies.',
 'Table 2.6 specifies the adapted questions. Sample selection covers representative sampling and exclusions. Index-test assessment covers an untouched evaluation set, separation of patients/isolates/plates/slides or acquisition runs before cropping and augmentation, training-only fitting of data-dependent preprocessing, and selection of models and thresholds without test-set feedback. Reference-standard assessment covers validity, independence from the index model and blinded interpretation where relevant. Flow and timing cover whether all eligible evaluation units received the intended reference, reasons for missing or excluded units, and alignment of the imaged growth stage with the reference observation.',
 'Applicable signalling questions will be answered Yes, No or Unclear, with supporting source locations. A domain is Low when its applicable questions are satisfactorily addressed; a documented material mechanism of bias warrants High. A No answer prompts a reasoned domain judgement rather than an automatic score. Unclear is reserved for insufficient reporting. The review-level overall rule is Low only when all four risk domains are Low, High when any is High, and Unclear otherwise. Applicability is rated separately for sample selection, index test and reference standard, using the same overall logic. No numerical quality score will be calculated, and tier, accuracy, dataset openness or absent external validation alone will not determine risk of bias [65].',
 'The planned assessment requires two human reviewers to pilot the adapted tool on a subset, assess eligible evaluations independently, and resolve differences by discussion or a third reviewer. Each reviewer and any adjudicator will be named in the assessment record. Independent human QUADAS-2 assessment has not yet been completed in the current review; the individual AI tiering pass is not a substitute. Supplementary Table S3 provides a 206-report worksheet with all judgement cells unfilled and assessment status recorded explicitly. Accordingly, this report presents the assessment framework without claiming completed domain ratings or reviewer agreement.'
]
replace_between('2.6. Risk of Bias and Quality Assessment','2.7. Data Synthesis',
    [normal(x) for x in quadas_paragraphs[:3]]+[('Table 2.6. QUADAS-2 signalling questions adapted to bacterial image analysis.','Caption')]+[normal(x) for x in quadas_paragraphs[3:]])
qrows=[
 ('D1 Sample selection','Were specimens, isolates or plates sampled consecutively, randomly or by a justified representative design? Were difficult, mixed, low-burden and poor-quality samples retained or exclusions justified? Was an artificial case-control contrast avoided?'),
 ('D2 Index test','Was evaluation independent of training, model selection and threshold tuning? Were source units separated before crops/augmentation? Were preprocessing and feature selection fitted only on training data? Was leakage or test-feedback avoided?'),
 ('D3 Reference standard','Does the reference establish the actual task target? Is it independent of the evaluated model? Were assessors blinded where interpretation is subjective? Were label adjudication and microbiological confirmation described?'),
 ('D4 Flow and timing','Did all eligible test units receive the intended reference and enter the analysis? Were failures and exclusions explained? Were reference observation and imaging comparable in time, specimen and growth stage?'),
 ('Applicability A1-A3','Do samples and setting, the index pipeline and the reference definition match the review question and intended use? Judge each of these three concerns separately; do not add a flow/timing applicability domain.')]
add_table(find('Table 2.6.')._p,['Domain','Adapted signalling questions'],qrows,[1.1,5.1])
find('Applicable signalling questions').paragraph_format.space_before=Pt(6)

qual=[
 'We organised the 206 reports in a study-level synthesis matrix by species or genera, dataset source and size, acquisition modality and conditions, explicit feature type, model family, task, reference standard and validation design (Supplementary Tables S1-S2). The 52 agar-plate reports were retained as a distinct subgroup. Taxa, tasks, modalities and reported acquisition factors may overlap; counts therefore refer to reports and are not assumed to sum to 206. Public-dataset reuse and repeated reports using potentially shared specimens were flagged when recognisable, because the number of papers is not the number of independent cohorts.',
 'For RQ2, we distinguished within-study comparisons of acquisition, preprocessing or augmentation from authors merely reporting that they used these procedures. We summarised changes in medium, incubation time, illumination, device, magnification, resolution, stain and image quality, together with preprocessing and synthetic-data strategies. Where the paper supplied a matched comparison, we recorded the baseline, changed condition, evaluation unit, metric and direction/magnitude of change. We kept within-domain tests separate from cross-device, cross-site and cross-condition tests. A list of factors mentioned by authors was not interpreted as a measured treatment effect.',
 'For RQ3, we separated explicit morphology, texture, topology or motility descriptors from features learned directly from images or spectral inputs, and identified fusion models separately. The clearest comparisons were those evaluating handcrafted features alone, image/spectral inputs alone and their combination on the same samples, split and metric. Between-paper differences were interpreted in the context of species, task, class balance, acquisition and validation. Features inferred from a still colony image were not treated as a direct motility measurement. In the absence of matched agar-plate comparisons, we retained a narrative synthesis rather than attributing cross-study accuracy differences to feature choice.'
]
replace_between('2.7.1. Qualitative synthesis','2.7.2. Quantitative synthesis',[normal(x) for x in qual])
quant=[
 'We summarised study characteristics and evidence-tier membership as n/N and percentages with two-sided Wilson 95% confidence intervals. The denominator was 206 for the complete review and 52 for the agar-plate subgroup; field-specific summaries identified their eligible denominator and missing values. These intervals describe proportions of reports, not patient-level diagnostic performance. For each report with an eligible numeric acc_pct value in the coded extraction, we used one existing headline accuracy value and reported the number contributing, median and 25th-75th percentiles, using linearly interpolated sample quantiles. No new classification-accuracy values were derived from mAP, precision, F1, Dice, counting agreement or relative counting error. Existing headline values follow the extraction rules in Section 2.5; this tiering pass did not independently re-audit all outcome metrics. The 126 available headline values remain heterogeneous in task and evaluation unit; their medians are descriptive summaries of reported results, not pooled diagnostic accuracy.',
 'Exploratory tier comparisons used tie-corrected Kruskal-Wallis statistics with 9,999 label permutations (seed 20261028), separately for all eligible headline values, agar-plate reports and other image reports. Holm adjustment covered these three omnibus comparisons; no post-hoc pairwise tests were performed. Such rank tests do not adjust for different tasks, case mix, reference standards, class numbers or data reuse. Their p-values are descriptive screens, with independence potentially violated by shared datasets, and were not interpreted as evidence that one model family or tier performs better. These analyses were specified during the present synthesis and were not a prospectively registered confirmatory analysis.',
 'We would consider bivariate random-effects diagnostic meta-analysis only for a clinically and methodologically coherent subset with paired sensitivity and specificity, reconstructable TP/FP/FN/TN counts at an independent specimen/isolate/plate level, comparable target states, reference standards and test thresholds, and enough independent cohorts for stable estimation [66]. A binomial-likelihood model with correlated random effects on logit sensitivity and specificity would preserve their relationship. Binary detection, multiclass species identification, pixel segmentation and colony counting would not be pooled together. Overall accuracy alone, or a multiclass confusion matrix without a defined common binary target, is insufficient. The current extraction has no harmonised cohort-level 2 x 2 dataset, and the full review spans incompatible tasks and sampling units; no pooled model was fitted.',
 'If a suitable diagnostic subset becomes available, heterogeneity will be examined jointly through paired forest plots, ROC-space patterns, between-study variances/covariance and prediction regions. I-squared and Cochran Q may be reported for a justified, separate univariate synthesis, but will not replace joint assessment of sensitivity and specificity or be calculated from the unweighted headline-accuracy distribution. Sensitivity analyses in this pass excluded low-confidence tier assignments, retained only high-confidence assignments, excluded R01154 (unverified final publication version) and R00395 (unresolved bacterial-outcome scope), and excluded reports carrying the AGAR/DIBaS reuse flags. These are report-level robustness checks, not proof of cohort independence. Exclusion of high-risk studies and leave-one-independent-cohort-out meta-analysis remain conditional on completed risk ratings and a suitable pooled dataset.',
 'Publication bias and selective reporting were considered qualitatively through unavailable outcomes, absent denominators and emphasis on selected best models. Formal diagnostic funnel asymmetry would use the Deeks effective-sample-size approach only within an eligible meta-analytic subset with sufficient independent studies (at least 10 as a pragmatic minimum), with cautious interpretation [67]. We did not apply Egger testing or trim-and-fill to heterogeneous headline accuracies. No meta-analysis, I-squared estimate or publication-bias test is reported here; their absence does not establish homogeneity or exclude publication bias. Tables, point data, tier decisions and Figure 5 are reproducible from the saved local scripts.'
]
replace_between('2.7.2. Quantitative synthesis','3. Results',[normal(x) for x in quant])

# New Results subsection, leaving existing characteristic and RQ text intact.
anchor=find('3.3. RQ1:')._p
prev=anchor.getprevious()
p=insert_after(prev,'3.2.3. Evidence tiers','Heading 3')
all_dist=dist[dist.subgroup=='All included'].set_index('tier')
def ci_text(t):
    r=all_dist.loc[t];return f'{int(r["n"])}/206 ({r.percent:.1f}%; 95% CI {r.ci_low_percent:.1f}-{r.ci_high_percent:.1f})'
text=(f'The source-based AI review assigned {ci_text("I")} to Tier I, {ci_text("II")} to Tier II and {ci_text("III")} to Tier III (Table 3.5; Figure 5). '
      'No report met the evidence requirement for Tier S (0/206; Wilson 95% CI 0.0-1.8%). This means that live prospective deployment was not established by the reviewed local sources; it is not evidence that no such system exists. '
      'Among the 52 agar-plate reports, 14 (26.9%) were Tier I, 29 (55.8%) Tier II and 9 (17.3%) Tier III. '
      'Prospective collection of clinical specimens, retrospective hospital data and the existence of an automated prototype were not equated with prospective live-workflow evaluation.')
p=insert_after(p._p,text)
p=insert_after(p._p,'Confidence was high for 161 assignments (78.2%), medium for 35 (17.0%) and low for 10 (4.9%). The 45 medium/low-confidence decisions are listed in a verification queue, including uncertainty about dataset ownership, routine-sample provenance and overlapping prototype/reference-strain rules. R01154 was assessed from its local 2025 preprint; the cited final journal version remains unverified. These assignments represent an individual AI review, with independent human verification still pending.')
p=insert_after(p._p,'Table 3.5. Evaluation-setting tiers and as-extracted headline accuracy. Accuracy n is the number of reports contributing a numeric value; medians are not pooled diagnostic effects.','Caption')
ar=acc[acc.cohort=='All included'].set_index('tier')
agar_dist=dist[dist.subgroup=='Agar plates'].set_index('tier')
rows=[]
for t in ['S','I','II','III']:
    r=all_dist.loc[t];a=ar.loc[t];g=agar_dist.loc[t]
    rows.append([t,f'{int(r["n"])} ({r.percent:.1f}%)',f'{r.ci_low_percent:.1f}-{r.ci_high_percent:.1f}',
                 f'{int(g["n"])} ({g.percent:.1f}%)',int(a.n_accuracy),
                 'Not available' if not a.n_accuracy else f'{rounded2(a["median"])} ({rounded2(a.q1)}-{rounded2(a.q3)})'])
el=add_table(p._p,['Tier','All n (%)','95% CI (%)','Agar n (%)','Accuracy n','Median accuracy % (IQR)'],rows,[.38,.93,1.03,1.0,.65,1.7])
p=insert_after(el,'')
p.paragraph_format.keep_with_next=True
p.add_run().add_picture(str(ROOT/'report_outputs/figures/evidence_tier_distribution.png'),width=Inches(6.1))
p=insert_after(p._p,'Figure 5. Distribution of evaluation-setting tiers among 206 reports, with Wilson 95% confidence intervals. Tier measures setting realism and is separate from risk of bias.','Caption')
p=insert_after(p._p,'Across the 126 reports with headline accuracy, the median was 93.66% (IQR 87.58-97.65; n = 56) in Tier I, 96.75% (90.73-98.60; n = 48) in Tier II and 98.18% (87.54-99.34; n = 22) in Tier III. The exploratory global rank comparison gave H = 5.20 and permutation p = 0.0756 (Holm-adjusted p = 0.1512). For agar-plate reports, only three Tier III papers contributed accuracy; their median was 87.16%, compared with 92.00% in Tier I and 97.50% in Tier II. The corresponding unadjusted p = 0.0398 became 0.1194 after Holm adjustment. Differences in metric definition, task and case mix prevent these results from establishing a deployment-related performance gradient.')
p=insert_after(p._p,'Excluding the 10 low-confidence assignments left 196 reports: Tier I 74 (37.8%), Tier II 93 (47.4%) and Tier III 29 (14.8%). Restricting to 161 high-confidence assignments gave 60 (37.3%), 75 (46.6%) and 26 (16.1%), respectively. Excluding R01154 and R00395 left 204 reports with 74 (36.3%), 98 (48.0%) and 32 (15.7%) in the three tiers. No sensitivity set contained a Tier S assignment. Full accuracy summaries and dataset-reuse sensitivity results are saved with the analysis tables; these checks do not resolve all shared-cohort dependencies.')

# Replace the empty RoB-results block with its actual status, without inventing ratings.
replace_between('3.6. Quality Assessment and Risk of Bias','4. Discussion',[
 normal('Independent human QUADAS-2 assessment is pending. Therefore, domain-level and overall Low/High/Unclear counts, reviewer agreement and risk-of-bias-based sensitivity analyses are not available. Table 2.6 and Supplementary Table S3 specify the assessment framework and worksheet. The source-based tier results in Section 3.2.3 cannot be interpreted as risk-of-bias ratings.'),
 normal('The source review identifies issues for that assessment, including selected strain panels or retrospective case series, possible dependence between image crops, test-set feedback during model development, references that verify counts or morphology rather than species identity, and exclusions of failed or poor-quality images. Public data, external validation or a high headline accuracy do not settle these domains. These observations guide the pending assessment and are not a substitute for supported per-study judgements.')])

# References append after the existing numbered study references, before Annex A.
annex=find('Annex A.')._p
last=annex.getprevious()
refs=[
 '[65] P. F. Whiting et al., “QUADAS-2: A revised tool for the quality assessment of diagnostic accuracy studies,” Ann. Intern. Med., vol. 155, no. 8, pp. 529-536, 2011, doi: 10.7326/0003-4819-155-8-201110180-00009.',
 '[66] J. J. Deeks, P. M. Bossuyt, M. M. Leeflang, and Y. Takwoingi, Eds., Cochrane Handbook for Systematic Reviews of Diagnostic Test Accuracy, version 2.0, 2023. doi: 10.1002/9781119756194.',
 '[67] J. J. Deeks, P. Macaskill, and L. Irwig, “The performance of tests of publication bias and other sample size effects in systematic reviews of diagnostic test accuracy was assessed,” J. Clin. Epidemiol., vol. 58, no. 9, pp. 882-893, 2005, doi: 10.1016/j.jclinepi.2005.01.016.'
]
for text in refs:last=insert_after(last,text)._p
replace(find('The repository holds everything needed to rebuild'),
 'The repository stores screening, extraction and analysis files. For the tier and synthesis update, 07_prepare_tier_review.py retrieves page-labelled source text without assigning tiers; manual_decisions.tsv records the individual source-based judgements. Run 08_summarise_evidence_tiers.py to rebuild the tier tables, sensitivity summaries, exploratory tests and Figure 5, then 09_update_tier_and_synthesis_report.py to reproduce this report update from its pre-edit backup. These scripts need Python 3, pandas, numpy, Pillow and python-docx; PDF-text retrieval additionally needs pypdf and pdfplumber. The legacy scripts remain available for earlier analyses, but 05_write_report.py uses an older source document and must not be run over this revised report. Supplementary Table S2 joins tier decisions to the synthesis matrix; S3 is the unfilled human QUADAS-2 assessment worksheet. The review package includes the complete figure data and code.')
replace(find('[Optional further annexes:'),
 'Supplementary Table S2: individual tier decisions, source locations, confidence and study synthesis matrix (206 reports). Supplementary Table S3: independent human QUADAS-2 assessment worksheet (not yet rated). Additional saved tables contain tier frequencies and confidence intervals, model-family cross-tabulations, accuracy point data, exploratory tests and sensitivity summaries. See results/final/tier_review/README.txt for scope and reproduction instructions.')
# Keep each new table header attached to its next row.
for node in d._element.iter(qn('w:t')):
    if node.text:
        node.text=node.text.replace('The risk-of-bias assessment appears in Section 3.6.',
            'Independent QUADAS-2 ratings remain pending (Section 3.6).')
for t in d.tables:
    if t.rows:
        for p in t.rows[0].cells[0].paragraphs:p.paragraph_format.keep_with_next=True
d.save(REPORT)
print(REPORT)
