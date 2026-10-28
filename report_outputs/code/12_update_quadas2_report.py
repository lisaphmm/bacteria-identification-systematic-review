"""Insert the completed single-AI QUADAS-2 appraisal into the current manuscript.

Reproduction starts from the immutable pre-QUADAS backup, preserving tier work.
Never rerun legacy builders 05/09 over the revised report without this step.
"""
import json,shutil
from pathlib import Path
from decimal import Decimal,ROUND_HALF_UP
import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parents[2]
REPORT=ROOT/'protocol/Systematic Review.docx'
BACKUP=ROOT/'protocol/Systematic Review (backup before QUADAS2 assessment 2026-10-28).docx'
if not BACKUP.exists():shutil.copy2(REPORT,BACKUP)
d=Document(BACKUP)
OUT=ROOT/'results/final/quadas2_review';TABLES=ROOT/'report_outputs/tables';FIGS=ROOT/'report_outputs/figures'
s=json.loads((OUT/'analysis_summary.json').read_text())
assert s['overall_risk']=={'U':122,'H':82,'L':2}
dist=pd.read_csv(TABLES/'quadas2_domain_distribution.csv')
acc=pd.read_csv(TABLES/'quadas2_accuracy_sensitivity.csv')
def find(prefix):return next(p for p in d.paragraphs if p.text.strip().startswith(prefix))
def replace(p,text):
 for c in list(p._p):
  if c.tag!=qn('w:pPr'):p._p.remove(c)
 p.add_run(text);return p
def after(el,text='',style='Normal'):
 n=OxmlElement('w:p');el.addnext(n);p=Paragraph(n,d._body);p.style=d.styles[style];p.add_run(text);return p
def clear_between(start,end):
 a=find(start)._p;b=find(end)._p;n=a.getnext()
 while n is not None and n is not b:
  nxt=n.getnext();n.getparent().remove(n);n=nxt
 return a
def table(el,headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
 for j,w in enumerate(widths):t.columns[j].width=Inches(w)
 for j,v in enumerate(headers):t.rows[0].cells[j].text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for i,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j])
   for p in c.paragraphs:
    p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
    for r in p.runs:r.font.size=Pt(8);r.bold=i==0
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 el.addnext(t._tbl);return t._tbl
def para(el,text,style='Normal'):return after(el,text,style)._p
def picture(el,name,caption,width=6.1,pagebreak=False):
 p=after(el);p.paragraph_format.keep_with_next=True
 p.paragraph_format.space_after=Pt(2)
 p.paragraph_format.page_break_before=pagebreak
 p.add_run().add_picture(str(FIGS/name),width=Inches(width))
 return after(p._p,caption,'Caption')._p
def rnd(v):return str(Decimal(str(v)).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP))

# Completed methods, with explicit reviewer identity and adaptation limits.
el=clear_between('2.6. Risk of Bias and Quality Assessment','2.7. Data Synthesis')
methods=[
 'On 28 October 2026, one AI reviewer (Codex) performed a separate source-based appraisal of all 206 included reports using a review-specific adaptation of QUADAS-2 [65]. The reviewer read relevant dataset, acquisition, validation, reference-standard and participant/image-flow passages in each local full-text PDF, following up ambiguous passages and visually checking conflicting source pages where needed. Existing extraction was used for navigation; prior extraction and tiering were available, so this was not a blinded review. Individual domain judgements and reasons were authored through reading, while code checked completeness, derived overall labels and produced tables and figures. No second reviewer, independent human validation, consensus adjudication or inter-reviewer agreement statistic was available.',
 'The sample unit was the patient specimen, isolate, plate, slide, original image, cultivation/acquisition run or time-lapse movie appropriate to the task. Colonies, cell crops and repeated frames nested within a source unit were not automatically treated as independent samples. For colony identification, the target was the organism identity at the stated taxonomic level; correct identification meant agreement with the reference identity. The index test was the complete AI/ML pipeline, including preprocessing, feature extraction, fitting, model selection and decision thresholds. Where a paper reported both internal development and a genuinely separate final evaluation, their evidential units and limitations were distinguished.',
 'Reference standards were matched to outcomes. MALDI-TOF, validated biochemical identification, sequencing or independently known culture identity could establish organism identity. Expert object annotations, masks and manual counts could establish locations, boundaries or numbers, but did not alone verify species. Conventional broth/agar dilution or another justified susceptibility method supported AST outcomes. Fluorescence probes, culture controls or manual trajectories were considered for their stated experimental targets. Morphology, microbial DNA presence, culture positivity and clinical infection were not assumed to be interchangeable reference constructs.',
 'Table 2.6 gives the adapted signalling questions. Eleven applicable answers per report were recorded as Yes, No or Unclear. Risk was judged in four domains and applicability separately in three. Low required satisfactory safeguards; a documented material bias mechanism supported High; insufficient reporting was Unclear. A No answer informed a reasoned domain judgement rather than a numerical score. Missing patient/plate grouping alone was not treated as proof of leakage; explicit random partitioning of correlated crops or augmented copies could support High. Public data, small samples, high accuracy, lack of code or absence of external validation alone did not determine a risk rating.',
]
for text in methods:el=para(el,text)
el=para(el,'Table 2.6. QUADAS-2 signalling questions adapted to bacterial image-analysis evaluations.','Caption')
qrows=[
 ['D1 Sample selection (3)','Representative source-unit sampling; avoidance of artificial disease case-control contrasts; avoidance or justification of exclusions.'],
 ['D2 Index test (2)','Evaluation without feedback into model/threshold selection; source-unit separation before crops/augmentation and training-only fitting of data-dependent preprocessing.'],
 ['D3 Reference standard (2)','Reference validity for the actual task; independence from the index model and blinded subjective interpretation where relevant.'],
 ['D4 Flow and timing (4)','Appropriate timing/growth stage; intended reference for all evaluated units; comparable reference procedures; all eligible units analysed or losses justified.'],
 ['Applicability A1-A3','Judge sample/spectrum, index task/pipeline and reference definition separately against the broad review and stated intended use. There is no flow/timing applicability domain.']]
el=table(el,['Domain / answer count','Adapted assessment'],qrows,[1.35,4.85])
el=para(el,'Overall risk was High if any risk domain was High, Low only if all four were Low, and Unclear otherwise. The same rule was applied separately to the three applicability judgements. No summed quality score was calculated. Supplementary Table S3 contains all 206 ratings, 1,442 domain-specific reasons, 2,266 signalling answers, source page/section locations, reviewer identity and local-PDF identifiers. Same-reviewer consistency refinements are recorded separately and are not inter-reviewer agreement. The 52 agar-plate reports were retained as a subgroup; unresolved scope and local-version concerns were flagged without silently changing the included denominator.')
el=para(el,'QUADAS-2 was developed for diagnostic accuracy. Its extension here to segmentation, counting, tracking and experimental phenotype/AST tasks is review-specific and has not been validated as a new measurement instrument. Applicability was assessed against the broad admitted bacterial-image task spectrum, not solely clinical agar identification: controlled cultures could be appropriate for an experimental segmentation or motility target, while selected strain panels could restrict broader clinical identification claims. The single-AI appraisal remains open to independent human verification [65].')

# Update synthesis to describe the analysis actually performed.
for p in d.paragraphs:
 if p.text.startswith('If a suitable diagnostic subset becomes available, heterogeneity'):
  replace(p,p.text.replace('Exclusion of high-risk studies and leave-one-independent-cohort-out meta-analysis remain conditional on completed risk ratings and a suitable pooled dataset.',
   'After the source-based QUADAS-2 appraisal, descriptive sensitivity analyses excluded High-risk reports and separately summarised Low and Unclear groups, overall and for agar plates. Excluding High leaves Low plus Unclear; it is not a restriction to low-bias evidence. Headline accuracy values were reused from extraction without a new outcome audit. Because only two reports were overall Low and the outcomes remain heterogeneous, no new inferential comparison by risk group was added. Leave-one-independent-cohort-out diagnostic meta-analysis remains conditional on a suitable harmonised dataset.'))
 if p.text.startswith('Publication bias and selective reporting were considered'):
  replace(p,p.text.replace('Tables, point data, tier decisions and Figure 5 are reproducible from the saved local scripts.','Tables, point data, tier decisions and QUADAS-2 figures are reproducible from the saved decision records and local scripts.'))

# Full results and their practical limits.
el=clear_between('3.6. Quality Assessment and Risk of Bias','4. Discussion')
el=para(el,'All 206 reports received a source-based single-AI appraisal (Table 3.9; Figure 8). Overall risk was High for 82 reports (39.8%; Wilson 95% CI 33.4-46.6%), Unclear for 122 (59.2%; 52.4-65.7%) and Low for 2 (1.0%; 0.3-3.5%). Unclear predominated because source sampling, biological-source separation and reference independence were often insufficiently reported. This does not mean that all Unclear reports had demonstrated bias. The overall rule requires all four domains to be Low, so the small Low group should not be read as a ranking of scientific value. Independent human confirmation and inter-reviewer reliability remain unavailable.')
el=para(el,'Table 3.9. Adapted QUADAS-2 judgements for 206 reports. Values are n (%); risk and applicability are separate constructs.','Caption')
domainnames={'D1':'D1 Sample selection','D2':'D2 Index test','D3':'D3 Reference standard','D4':'D4 Flow and timing','overall_risk':'Overall risk','A1':'A1 Sample applicability','A2':'A2 Index applicability','A3':'A3 Reference applicability','overall_applicability':'Overall applicability'}
rows=[]
for key,label in domainnames.items():
 g=dist[(dist.cohort=='All included')&(dist.domain==key)].set_index('judgement')
 rows.append([label]+[f'{int(g.loc[c,"n"])} ({g.loc[c,"percent"]:.1f}%)' for c in ['Low','High','Unclear']])
el=table(el,['Domain','Low','High','Unclear'],rows,[2.15,1.35,1.35,1.35])
el=picture(el,'quadas2_domain_summary.png','Figure 8. Distribution of source-based adapted QUADAS-2 risk and applicability judgements. Narrow segments show counts only; full percentages and Wilson intervals are in the accompanying table data. One AI reviewer; no numerical quality score.')
el=para(el,'Sample selection was High in 35 reports (17.0%) and Unclear in 167 (81.1%). Supported concerns included removal of reference-discordant samples, difficult early/mixed cultures, poor-quality images or unresolved expert labels. Deliberate laboratory class panels were not automatically classified as a disease case-control design. Background-only patch rejection and valid assay quality-control criteria were distinguished from exclusion of diagnostically difficult eligible samples; uncertain recruitment was retained as Unclear.')
el=para(el,'Index-test risk was High in 45 reports (21.8%), Low in 53 (25.7%) and Unclear in 108 (52.4%). Explicit test-feedback examples included stopping training only after test accuracy exceeded a target (R01980) and early stopping based on test accuracy (R02940). Other concerns involved random partitioning of source-correlated patches or augmented copies. A missing grouping statement alone was classified as Unclear. Stronger designs included an independent 76-strain evaluation in R01626, experiment-level holdouts in R01853, and separate development, calibration and specimen-level testing where reported. Internal image metrics were not automatically promoted to patient-level evidence.')
el=para(el,'Reference-standard risk was High in 10 reports (4.9%) and Unclear in 128 (62.1%). Principal concerns included evaluating biological segmentation against unconfirmed algorithm-generated masks, assigning clinical urinary-infection labels from organism categories, and revising reference diagnoses after inspecting index disagreements. Expert masks can be suitable for segmentation while remaining insufficient for species identity. Independent culture, molecular identity or susceptibility measurements supported task-matched Low judgements where their use and independence were clear. Human readers used only as performance comparators were not misidentified as reference-standard reviewers.')
el=para(el,'Flow and timing was High in 31 reports (15.0%). R03920 excluded 34 scanning failures from 320 slides before accuracy analysis. R07472 reported its headline whole-slide classification metrics among 160 slides where bacteria were detected, while 29 background-classified slides fell outside those conditional metrics. These endpoints do not describe all-attempt workflow performance. Conversely, R05464 distinguished 263 clinical culture validations from the 215-sample paired-urinalysis comparator subset; missing comparator urinalysis was not treated as exclusion from the primary culture evaluation. R05350 had conflicting validation counts in Methods and Results and unconfirmed reference-positive cases; its pages 4-5 were visually checked and the discrepancy retained.')
el=para(el,'Overall applicability concern was High in 94 reports (45.6%), Low in 64 (31.1%) and Unclear in 48 (23.3%). Sample/spectrum concerns predominated (87 High), while index-task applicability was Low in 205 reports under the broad included scope. Reference-definition concerns were High in 13 reports. R00395 retained an unresolved principal nonbacterial-outcome scope flag; R01154 was assessed from the local preprint without claiming verification of the cited final journal version. Such flags require scope/version adjudication and were not used to silently delete papers or change the tier assignments.')
el=para(el,'Within the 52 agar-plate reports, overall risk was High in 16 (30.8%) and Unclear in 36 (69.2%); none was overall Low. Among the other 154 reports, 66 were High, 86 Unclear and 2 Low. These report-level distributions do not measure independent cohort counts because public datasets and related source material recur across publications. The complete per-study traffic-light display appears in Supplementary Figure S1a-d (Annex B), with all reasons and source locations in Supplementary Table S3.')
el=para(el,'Table 3.10. Descriptive risk-based sensitivity of as-extracted headline accuracy. IQR is the 25th-75th percentile; values are not pooled diagnostic accuracy.','Caption')
rows=[]
for _,r in acc.iterrows():
 if r.cohort=='Agar plates' and r.analysis not in ['All reports','Exclude High risk (Low + Unclear)']:continue
 label=('All studies: ' if r.cohort=='All included' else 'Agar: ')+r.analysis
 value='Not available' if r.accuracy_n==0 else f'{rnd(r["median"])} ({rnd(r.q1)}-{rnd(r.q3)})'
 rows.append([label,int(r.reports),int(r.accuracy_n),value])
el=table(el,['Analysis subset','Reports','Accuracy n','Median % (IQR)'],rows,[2.8,.65,.75,2.0])
el=para(el,'Across 126 reports with an extracted numeric headline accuracy, the median was 95.03% (IQR 88.88-98.52%). Excluding High-risk reports left 124 reports, of which 74 had accuracy values: median 95.25% (90.15-98.43%). The High-risk group contributed 52 values: 94.98% (87.98-98.54%). The remaining 122 Unclear reports supplied 72 values; the two overall Low reports supplied only two. For agar plates, 29 values gave 96.50% (90.97-98.52%); excluding High left 20 values with 97.05% (91.43-98.85%). Figure 9 shows the individual reported values. The numerical similarity does not validate the algorithms or negate bias: tasks, definitions, source cohorts and evaluation units differ, and excluding High predominantly retains Unclear evidence.')
el=picture(el,'quadas2_accuracy_sensitivity.png','Figure 9. As-extracted headline accuracy by overall adapted QUADAS-2 risk. Each point is one report, with IQR and median overlays. Two Low-risk reports are insufficient for a stable low-risk benchmark; no inferential risk-group test or pooled diagnostic model was fitted.')
el=para(el,'No bivariate random-effects meta-analysis, I-squared estimate or diagnostic publication-bias test was added. The appraisal does not supply harmonised independent-cohort TP/FP/FN/TN data or resolve incompatible targets and thresholds. Likewise, no agreement statistic was calculated from a single reviewer. The separate workbook and CSV exports provide the complete decisions; all chart inputs and figure-generation code are saved alongside the report.')

# Remove stale pending-rating claims without inventing completed human review.
for p in d.paragraphs:
 if p.text.startswith('Methodological quality and reporting.'):
  replace(p,'Methodological quality and reporting. The single-AI QUADAS-2 appraisal found 82 High, 122 Unclear and 2 Low overall-risk reports. Missing source-sampling and validation/reference-independence details were common. Excluding High left mostly Unclear evidence and changed the descriptive accuracy median little; it did not establish a robust low-risk evidence base. Independent human verification remains pending (Section 3.6).')
 if p.text.startswith('The repository stores screening, extraction and analysis files.'):
  replace(p,'The repository stores screening, extraction and analysis records. Tier decisions and their analyses remain in results/final/tier_review. The new QUADAS-2 decision record, assessment protocol, consistency audit, full study/domain/question exports and separate workbook are in results/final/quadas2_review. Script 10_read_quadas_sources.py retrieves page-labelled navigation passages without assigning ratings; 11_summarise_quadas2.py aggregates manually authored decisions and generates figures; 11b_build_quadas2_workbook.mjs exports the workbook; 12_update_quadas2_report.py updates this manuscript from its pre-QUADAS backup. Earlier report builders use older snapshots and must not overwrite this version without reapplying the later updates. The human-review worksheet is a separate unfilled template; it is not the completed AI assessment.')
 if p.text.startswith('Supplementary Table S2: individual tier decisions'):
  replace(p,'Supplementary Table S2 records individual evidence tiers and the study synthesis matrix (206 reports). Supplementary Table S3 records the completed single-AI QUADAS-2 appraisal: 206 study-level rows, 1442 domain-specific evidence records and 2266 signalling answers, plus source locations and version/scope notes. The separate workbook contains Summary, Study ratings, Domain evidence, Signalling answers and Protocol sheets. CSV exports contain complete figure data, domain distributions with Wilson confidence intervals and accuracy sensitivity summaries. Independent human verification is not recorded as complete.')

# Full study display is placed in a distinct annex; main Results refers forward.
p=d.add_paragraph('Annex B. Study-level QUADAS-2 ratings',style='Heading 1');p.paragraph_format.page_break_before=True
el=d.paragraphs[-1]._p
for j,letter in enumerate('abcd',1):
 el=picture(el,f'quadas2_traffic_light_{j}.png',f'Supplementary Figure S1{letter}. Study-level adapted QUADAS-2 ratings, panel {j} of 4. Domain definitions and complete supporting evidence are in Supplementary Table S3.',width=5.9,pagebreak=j>1)
for t in d.tables:
 for row in t.rows:
  for cell in row.cells:
   if cell.text=='Figures 1 to 6 as PNG':
    cell.text='Figures 1 to 9 and Supplementary Figure S1a-d as PNG'
    for para in cell.paragraphs:
     for run in para.runs:run.font.size=Pt(9)
 if t.rows:
  for p in t.rows[0].cells[0].paragraphs:p.paragraph_format.keep_with_next=True
d.core_properties.comments='Updated with single-AI source-based adapted QUADAS-2 appraisal, 28 October 2026.'
d.save(REPORT)
print(REPORT)
